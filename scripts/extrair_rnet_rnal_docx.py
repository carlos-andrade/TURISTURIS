#!/usr/bin/env python3
"""Extrai de forma reprodutível todo o conteúdo relevante de DADOS/RNET_RNAL.docx."""
from __future__ import annotations
import csv, hashlib, json, re, shutil, zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "DADOS" / "RNET_RNAL.docx"
OUT = ROOT / "DADOS" / "RNET_RNAL"
RAW = OUT / "RAW_DOCX"
MEDIA = OUT / "MEDIA"
OUT.mkdir(parents=True, exist_ok=True)
RAW.mkdir(parents=True, exist_ok=True)
MEDIA.mkdir(parents=True, exist_ok=True)

if not SRC.exists():
    raise FileNotFoundError(SRC)

data = SRC.read_bytes()
sha = hashlib.sha256(data).hexdigest()

# Preserve the complete OOXML package, including XML and embedded media.
with zipfile.ZipFile(SRC) as z:
    members = []
    for info in z.infolist():
        members.append({"name": info.filename, "size": info.file_size, "compressed_size": info.compress_size})
        target = RAW / info.filename
        target.parent.mkdir(parents=True, exist_ok=True)
        if not info.is_dir():
            target.write_bytes(z.read(info.filename))
        if info.filename.startswith("word/media/") and not info.is_dir():
            out = MEDIA / Path(info.filename).name
            out.write_bytes(z.read(info.filename))

doc = Document(SRC)
paragraphs = []
for i, p in enumerate(doc.paragraphs, 1):
    text = p.text
    paragraphs.append({"index": i, "style": p.style.name if p.style else None, "text": text})

tables = []
for ti, table in enumerate(doc.tables, 1):
    rows = []
    for ri, row in enumerate(table.rows, 1):
        rows.append({"row": ri, "cells": [cell.text for cell in row.cells]})
    tables.append({"index": ti, "rows": rows})

full_text = "\n".join(p["text"] for p in paragraphs)
(OUT / "TEXTO_COMPLETO.txt").write_text(full_text + "\n", encoding="utf-8")
(OUT / "PARAGRAFOS.json").write_text(json.dumps(paragraphs, ensure_ascii=False, indent=2), encoding="utf-8")
(OUT / "TABELAS.json").write_text(json.dumps(tables, ensure_ascii=False, indent=2), encoding="utf-8")

with (OUT / "TABELAS.csv").open("w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(["tabela", "linha", "celula", "texto"])
    for t in tables:
        for row in t["rows"]:
            for ci, cell in enumerate(row["cells"], 1):
                w.writerow([t["index"], row["row"], ci, cell])

core = {}
core_xml = RAW / "docProps" / "core.xml"
if core_xml.exists():
    root = ET.fromstring(core_xml.read_bytes())
    ns = {"dc":"http://purl.org/dc/elements/1.1/", "dcterms":"http://purl.org/dc/terms/", "cp":"http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
    for key, path in {
        "title":"dc:title","subject":"dc:subject","creator":"dc:creator",
        "description":"dc:description","lastModifiedBy":"cp:lastModifiedBy",
        "created":"dcterms:created","modified":"dcterms:modified",
    }.items():
        node = root.find(path, ns)
        if node is not None:
            core[key] = node.text

manifest = {
    "schema_version": "1.0",
    "process": "TURISTURIS_EXTRACAO_RNET_RNAL_DOCX",
    "source": "DADOS/RNET_RNAL.docx",
    "capture_utc": datetime.now(timezone.utc).isoformat(),
    "source_sha256": sha,
    "source_size_bytes": len(data),
    "paragraph_count": len(paragraphs),
    "table_count": len(tables),
    "ooxml_member_count": len(members),
    "embedded_media_count": len(list(MEDIA.iterdir())),
    "core_properties": core,
    "ooxml_members": members,
    "outputs": [
        "TEXTO_COMPLETO.txt","PARAGRAFOS.json","TABELAS.json","TABELAS.csv",
        "RAW_DOCX/","MEDIA/","MANIFESTO_EXTRACAO.json"
    ],
}
(OUT / "MANIFESTO_EXTRACAO.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"result":"SUCCESS","source_sha256":sha,"paragraph_count":len(paragraphs),"table_count":len(tables),"media_count":manifest["embedded_media_count"]}, ensure_ascii=False))
